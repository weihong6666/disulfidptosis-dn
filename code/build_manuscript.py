# -*- coding: utf-8 -*-
"""
从 manuscript/manuscript_v3.0.md 生成投稿用 DOCX（并转 PDF）。

设计要点：
  * v3.0 Markdown 是唯一数据源，不再有第二份硬编码正文（旧版 生成投稿文档.py 的病根）
  * 图表在「Figure Legends」小节内、图注之后插入
  * 正文段落支持 **粗体** 与 *斜体*

输出：
  manuscript/manuscript_v3.0.docx
  manuscript/manuscript_v3.0.pdf   （经 Word COM 转换）
"""

import os
import re
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor

ROOT = r"C:\Users\weihong\Desktop\生物竞赛"
SRC = os.path.join(ROOT, "manuscript", "manuscript_v3.0.md")
OUT_DOCX = os.path.join(ROOT, "manuscript", "manuscript_v3.0.docx")
OUT_PDF = os.path.join(ROOT, "manuscript", "manuscript_v3.0.pdf")

# 图号 -> 图片文件（与 v3.0 的 Figure 1-10 一一对应）
FIGURES = {
    1: [("Fig1_study_design_journal.png", 6.0)],
    2: [("02_volcano.png", 5.2), ("02_DRDEGs_heatmap.png", 5.0)],
    3: [("03_soft_threshold.png", 4.6), ("03_module_trait_heatmap.png", 5.4)],
    4: [("04_lasso_cv.png", 4.8), ("04_svm_rfe_rank.png", 4.8), ("04_rf_importance.png", 4.8)],
    5: [("04_ROC_curves.png", 5.2), ("05_validation_ROC.png", 5.2)],
    6: [("06_nomogram.png", 4.8), ("06_calibration.png", 4.8), ("06_DCA.png", 4.8)],
    7: [("07_immune_boxplot.png", 5.6), ("07_ssGSEA_heatmap.png", 5.6),
        ("07_checkpoint_heatmap.png", 5.0), ("07_ESTIMATE_scores.png", 5.0)],
    8: [("02b_cell_death_boxplot.png", 6.0), ("02b_pathway_effect_size.png", 5.4),
        ("02b_oxidative_stress_heatmap.png", 5.0), ("02b_disulfidptosis_specificity.png", 4.8)],
    9: [("08_UMAP_clusters.png", 5.4), ("08_UMAP_group.png", 5.4),
        ("08c_biomarker_celltype_heatmap.png", 6.2), ("08b_pathway_violin.png", 5.2),
        ("08b_pseudobulk_logFC_heatmap.png", 5.2)],
    10: [("HPA_validation_heatmap.png", 5.4)],
}

SUPP = {
    "S1": [("04c_permutation_test.png", 5.0), ("04c_cv_auc_distribution.png", 5.0),
           ("04c_model_aic_comparison.png", 5.0)],
    "S2": [("04b_method_comparison.png", 5.4)],
    "S3": [("11_GO_BP_all_DEGs.png", 5.4), ("11_KEGG_all_DEGs.png", 5.4)],
    "S4": [("09_drug_target_network.png", 6.0), ("09_drug_gene_heatmap.png", 5.0)],
}

PAGE_BREAK_BEFORE = {
    "## 1. Introduction", "## 2. Methods", "## 3. Results", "## 4. Discussion",
    "## References", "## Figure Legends", "## Supplementary Materials",
}


def add_runs(par, text):
    """把 **粗体** 与 *斜体* 渲染成 Word run"""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def add_figure(doc, fname, width_in, missing):
    path = os.path.join(ROOT, fname)
    if not os.path.exists(path):
        missing.append(fname)
        return False
    par = doc.add_paragraph()
    par.alignment = 1  # center
    par.add_run().add_picture(path, width=Inches(width_in))
    return True


def build():
    md = open(SRC, encoding="utf-8").read()
    lines = md.split("\n")

    doc = Document()
    # 基础样式
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    missing = []
    in_fig_legends = False
    in_refs = False
    i = 0
    n_para = n_fig = 0

    while i < len(lines):
        raw = lines[i]
        s = raw.strip()

        # 空行
        if not s:
            i += 1
            continue

        # 水平线
        if s == "---":
            i += 1
            continue

        # 表格：连续的 | 行
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [
                [c.strip() for c in r.strip("|").split("|")]
                for r in block
                if not re.match(r"^\|[\s:\-|]+\|$", r)
            ]
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell in enumerate(row):
                        if c_idx >= len(rows[0]):
                            continue
                        par = t.cell(r_idx, c_idx).paragraphs[0]
                        add_runs(par, cell)
                        for run in par.runs:
                            run.font.size = Pt(9)
                            if r_idx == 0:
                                run.bold = True
                doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if s in PAGE_BREAK_BEFORE and (n_para or n_fig):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if level == 1:
                doc.add_heading(text, level=0)
            else:
                doc.add_heading(text, level=level - 1)
            if level == 2:
                in_fig_legends = s == "## Figure Legends"
                in_refs = s == "## References"
            i += 1
            continue

        # 列表项
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, s[2:])
            msup = re.match(r"^- \*\*Supplementary Figure (S\d)\.", s)
            if msup:
                for fname, w in SUPP.get(msup.group(1), []):
                    if add_figure(doc, fname, w, missing):
                        n_fig += 1
                doc.add_paragraph()
            i += 1
            continue

        # 图注 / 普通段落
        p = doc.add_paragraph()
        if in_refs:
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, s)
            for run in p.runs:
                run.font.size = Pt(9)
        else:
            add_runs(p, s)
            p.paragraph_format.space_after = Pt(8)
        n_para += 1

        # 图注后插入对应图片
        if in_fig_legends:
            mfig = re.match(r"^\*\*Figure (\d+)\.", s)
            if mfig:
                num = int(mfig.group(1))
                for fname, w in FIGURES.get(num, []):
                    if add_figure(doc, fname, w, missing):
                        n_fig += 1
                doc.add_paragraph()

        i += 1

    doc.save(OUT_DOCX)
    print("DOCX: %s" % OUT_DOCX)
    print("  段落 %d，嵌入图片 %d 张" % (n_para, n_fig))
    if missing:
        print("  !! 缺失图片 %d 张: %s" % (len(missing), missing))
    return missing


def to_pdf():
    try:
        from docx2pdf import convert
    except Exception as e:
        print("docx2pdf 不可用:", e)
        return False
    convert(OUT_DOCX, OUT_PDF)
    print("PDF : %s (%d B)" % (OUT_PDF, os.path.getsize(OUT_PDF)))
    return True


if __name__ == "__main__":
    missing = build()
    if "--no-pdf" not in sys.argv:
        to_pdf()
    sys.exit(1 if missing else 0)
